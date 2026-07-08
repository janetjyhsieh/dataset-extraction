from __future__ import annotations

import logging
from pathlib import Path
import openpyxl

from dataset_extraction.corpus.corpus import CorpusSource, SourcePaper
from dataset_extraction.corpus.utils import keywords_in_string

logger = logging.getLogger("dataset_extraction.corpus.xlsx")

_ANNOTATIONS_PATH = Path(__file__).parents[4] / "gdrive" / "annotations.xlsx"
_SHEET_NAME = "2. Papers"
_HEADERS_ROW = 2

_HTTP_HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; dataset-extraction-bot/1.0)"}

def _is_formula_error(value: str | None) -> bool:
    return not value or value.startswith("#") or value == "NA"

class XslxSource(CorpusSource):
    def __init__(self, working_dir, index_path, xsls_path):
        super().__init__(working_dir, index_path)
        self.xsls_path = xsls_path

    def get_papers(self, paper_sources, keywords=None):
        wb = openpyxl.load_workbook(str(self.xsls_path), data_only=True)
        ws = wb[_SHEET_NAME]

        headers = [cell.value for cell in ws[_HEADERS_ROW]]
        col = {h: i for i, h in enumerate(headers) if h}
        
        new_papers = []
        for row in ws.iter_rows(min_row=_HEADERS_ROW + 1, values_only=True):
            source = row[col["paper_source"]]
            if not source or source not in paper_sources:
                continue

            if row[col["STATUS"]] != "DONE":
                continue

            paper_id = row[col["paper_id"]]
            if not paper_id or paper_id in self.paper_ids:
                continue

            title = row[col["title"]]
            if not title:
                continue
            if keywords is not None and not keywords_in_string(title, keywords):
                continue

            url = row[col["url"]]
            url = None if _is_formula_error(url) else url

            new_papers.append(
                SourcePaper(
                    paper_id = paper_id,
                    title = title,
                    authors = [],
                    pdf_url = url,
                )
            )
            self.paper_ids.add(paper_id)
        self.papers.extend(new_papers)
        return new_papers
